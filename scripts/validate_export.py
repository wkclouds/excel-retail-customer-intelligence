"""Read-only verification of the delivered XLSX. Never writes the workbook."""
from pathlib import Path
import json, zipfile, xml.etree.ElementTree as ET
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
path=ROOT/'Retail_Customer_Intelligence.xlsx'
a=json.loads((ROOT/'data/processed/analysis.json').read_text(encoding='utf-8'))
values=openpyxl.load_workbook(path,data_only=True,read_only=True)
book=openpyxl.load_workbook(path,data_only=False,read_only=False)
expected=['Start_Here','Sales_Overview','Product_Analysis','Customer_RFM','Cohort_Retention','Regional_Analysis','Insights','Data_Quality','Data_Dictionary','Monthly_Data']
assert book.sheetnames==expected
checks=[]
for cell,metric in [('B10','gross'),('F10','returns'),('J10','net'),('B14','orders'),('F14','aov'),('J14','return_rate')]:
    actual=values['Sales_Overview'][cell].value
    assert isinstance(actual,(int,float)) and abs(actual-a['totals'][metric])<.00001,(cell,actual)
    checks.append({'cell':'Sales_Overview!'+cell,'metric':metric,'cached_value':actual,'passed':True})
errors=[]
for sheet in values:
    for row in sheet:
        for cell in row:
            if cell.data_type=='e' or (isinstance(cell.value,str) and 'is not implemented' in cell.value):
                errors.append(f'{sheet.title}!{cell.coordinate}: {cell.value}')
assert not errors,errors
assert all(values['Data_Quality'][f'E{r}'].value==0 for r in range(26,34))
assert values['Customer_RFM']['B17'].value=='Champions'
assert values['Sales_Overview']['H7'].value=='All countries'
assert values['Sales_Overview']['B7'].value=='2009-12'
assert values['Sales_Overview']['E7'].value=='2011-12'
assert len(book['Sales_Overview'].data_validations.dataValidation)==3
assert len(book['Customer_RFM'].data_validations.dataValidation)==1
assert sum(len(s._charts) for s in book)==4
assert len(book['Customer_RFM'].tables)==2
assert book['Customer_RFM'].tables['CustomerDetail'].ref==f'B38:L{38+len(a["rfm"])}'
assert book['Product_Analysis'].tables['ProductDetail'].ref==f'B33:J{33+len(a["products"])}'
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
with zipfile.ZipFile(path) as z:
    assert z.testzip() is None
    native_charts=[p for p in z.namelist() if '/charts/chart' in p and p.endswith('.xml')]
    assert len(native_charts)==4
    bindings=[]
    for p in native_charts:
        xml=ET.fromstring(z.read(p))
        refs=[e.text for e in xml.findall('.//c:f',ns)]
        assert len(refs)>=2,(p,refs)
        bindings.append({'chart':p,'source_references':refs})
    assert not any('vbaProject' in p or p.startswith('xl/externalLinks/') for p in z.namelist())
    assert not any(p.startswith('xl/pivotTables/') for p in z.namelist())
report={'read_only_parser':'openpyxl '+openpyxl.__version__,'xlsx_integrity':'passed','sheet_count':len(book.sheetnames),'native_chart_count':len(native_charts),'table_count':sum(len(s.tables) for s in book),'dropdown_count':sum(len(s.data_validations.dataValidation) for s in book),'formula_errors':errors,'cached_kpis':checks,'chart_bindings':bindings,'no_macros_or_external_workbook_links':True,'native_excel_opening':'Not performed','all_checks_passed':True}
(ROOT/'docs/export_validation.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps(report,indent=2))
